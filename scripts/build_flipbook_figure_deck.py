#!/usr/bin/env python3
"""Build annotated HTML + PowerPoint figure decks for flipbook red-pen review."""

from __future__ import annotations

import base64
import html
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt

REPO = Path(__file__).resolve().parents[1]
OUT_HTML = REPO / "3-Master_Plan/re_entry/_DISPOSABLE_paper_flipbook_PD30_FIGURE_DECK.html"
OUT_PPTX = REPO / "3-Master_Plan/re_entry/_DISPOSABLE_paper_flipbook_PD30_FIGURE_DECK.pptx"
OUT_INV_HTML = REPO / "3-Master_Plan/re_entry/_DISPOSABLE_paper_flipbook_PD30_FIGURE_INVENTORY.html"
OUT_INV_PPTX = REPO / "3-Master_Plan/re_entry/_DISPOSABLE_paper_flipbook_PD30_FIGURE_INVENTORY.pptx"
RE_ENTRY = REPO / "3-Master_Plan/re_entry"
ARMY_SANDBOX = REPO / "3-Master_Plan/re_entry/HEROs_and_PASSes/army_sandbox"

# Drop Cell-11 / Cell-12 exports here (preferred) or under talent/**/cox/cox_plots|cox_results.
ARMY_CANONICAL_FILENAMES: dict[str, tuple[str, ...]] = {
    "hero_cif": (
        "ARMY_G1_pool_minus_mean_fwd_cif_b8.png",
        "G1_pool_minus_mean_fwd_cif_b8.png",
    ),
    "own_tb_cif": (
        "ARMY_own_tb_ratio_fwd_cif_b8.png",
        "G1_own_tb_ratio_fwd_cif_b8.png",
    ),
    "partial_effects": (
        "ARMY_partial_effects_pool_minus_mean.png",
        "ARMY_cox_partial_effects.png",
    ),
}

FIGURE_ARMY_SLOTS: dict[tuple[str, str], str] = {
    ("Act 0", "0.3"): "hero_cif",
    ("Act I", "1.2"): "own_tb_cif",
    ("Act I", "1.3"): "hero_cif",
    ("Act I", "1.4"): "partial_effects",
}

INVENTORY_ARMY_SLOTS: dict[str, str] = {
    "Army hero porch (G1)": "hero_cif",
    "Army own-TB baseline": "own_tb_cif",
    "Army Cox / partial effects": "partial_effects",
}

# Inventory row: label, repo_path (None if no file), path_display, use_in_talk, note

SLIDE_W_IN = 13.333
SLIDE_H_IN = 7.5
MARGIN_IN = 0.4
HEADER_IN = 0.95  # compact caption strip — image gets the rest
FOOTER_IN = 0.55

SLIDE_W = Inches(SLIDE_W_IN)
SLIDE_H = Inches(SLIDE_H_IN)
MARGIN = Inches(MARGIN_IN)


def rel(p: Path) -> str:
    try:
        return p.relative_to(RE_ENTRY).as_posix()
    except ValueError:
        return p.relative_to(REPO).as_posix()


def status(path: Path | None) -> tuple[str, str]:
    if path is None:
        return "text", "TEXT — no figure file"
    if not path.exists():
        return "missing", f"MISSING — {rel(path)}"
    suf = path.suffix.lower()
    if suf == ".pdf":
        return "pdf", f"PDF — open locally: {rel(path)}"
    if suf in {".md", ".ipynb"}:
        return "text", f"DOC — {rel(path)}"
    return "ok", rel(path)


# (act, slide_id, title, use_line, path_from_repo_root or None, extra_note)
FIGURES: list[tuple[str, str, str, str, str | None, str]] = [
    # Act 0
    ("Act 0", "0.2", "What is *not* the claim?", "Binding three-separations table", "3-Master_Plan/BINDING_Selection_is_its_own_step.md", "Markdown — not a PNG; read binding doc."),
    ("Act 0", "0.3", "What is the hero?", "Army CIF / pool-quality porch (G1)", None, "G1: regenerate from talent/520 Cell 11 — z_pool_minus_mean_snr_fwd 8-bin promotion CIF."),
    ("Act 0", "0.4", "Why basketball in the repo?", "Access narrative", None, "Text slide — no figure."),
    # Act I
    ("Act I", "1.1", "Who is in the pond?", "Army cohort / HRC talent pipeline", "talent/documents/520_PIPELINE_COX_OVERVIEW.md", "Notebook: talent/talent_pipeline/520_pipeline_cox_working.ipynb"),
    ("Act I", "1.2", "Talent alone?", "Own TB ratio — monotone promotion CIF", None, "G1 sibling: Cell 11 z_tb_ratio_fwd_snr 8-bin bars (not in git)."),
    ("Act I", "1.3", "The hero shape?", "Strong inverted-U vs pool LOO (G1)", None, "G1 lead figure — Q1–Q8 promotion CIF on z_pool_minus_mean_snr_fwd."),
    ("Act I", "1.4", "Formal check?", "Cox quadratic + partial effects", None, "Cell 12 / 12.6A exports not in git."),
    # Act II — composites first in section via extras below
    ("Act II", "2.1", "Who is in the pond?", "MBB cohort (panel 1 text in 3×3)", "3-Master_Plan/re_entry/HEROs_and_PASSes/sports_sandbox/data_story/MBB_DATA_STORY_reigning_3x3.png", "Panel 1 = text tile top-left of this mosaic."),
    ("Act II", "2.2", "Talent alone? (naïve)", "Draft rate vs own ability ventiles", "3-Master_Plan/re_entry/HEROs_and_PASSes/pass_a/PASS_A_naive_ability_q16_allt_min20_mg10_9_21_last_ps.png", "Pairs with 2.3 hero; regen: pass_a_empirical_bundle.py --panel-rows last-ps --season-min 2009 --season-max 2021 --naive-panel"),
    ("Act II", "2.3", "Same axis, same family?", "Reigning HERO EW16 (panel 9)", "3-Master_Plan/re_entry/HEROs_and_PASSes/sports_sandbox/hero/HERO_ew16_allt_min20_mg10_09_21_last_ps_perm_loo_ever_lastps_ew16.png", ""),
    ("Act II", "2.4", "Assortment / overlap?", "Team interval overlap (panel 5)", "3-Master_Plan/re_entry/HEROs_and_PASSes/sports_sandbox/reigning_hero/basic_data_plots/REIGNING_team_interval_overlap_mg10_min20_09_21.png", ""),
    # Act III
    ("Act III", "3.1", "Same skeleton, academia?", "Tenure PD29 3×3 + panel 9 HERO", "3-Master_Plan/re_entry/HEROs_and_PASSes/tenure_sandbox/data_story/TENURE_DATA_STORY_pd29_3x3.png", "Panel 9 path below as standalone."),
    ("Act III", "3.1b", "Tenure HERO (panel 9 alone)", "Dept LOO career rate Q16", "3-Master_Plan/re_entry/HEROs_and_PASSes/tenure_sandbox/hero/HERO_tenure_q16_decision_dept_loo_infHM_slide.png", "Same as panel 9 in 3×3."),
    ("Act III", "3.2", "Which performance metric?", "Perf story p1", "3-Master_Plan/re_entry/HEROs_and_PASSes/tenure_sandbox/data_story/TENURE_PERF_METRIC_STORY.png", ""),
    ("Act III", "3.2b", "Perf story p2", "10-bin sensitivity", "3-Master_Plan/re_entry/HEROs_and_PASSes/tenure_sandbox/data_story/TENURE_PERF_METRIC_STORY_p2.png", ""),
    ("Act III", "3.3", "Three-leg composite?", "Army G1 + MBB + tenure heroes", "3-Master_Plan/re_entry/HEROs_and_PASSes/composites/THREE_LEG_HERO_composite.png", "Regen: scripts/build_three_leg_hero_composite.py"),
    # Act IV
    ("Act IV", "4.0", "Model in plain English", "Text intro before Model slide", None, "Text slide — four-step pipeline overview."),
    ("Act IV", "4.1", "Minimal generative story?", "Model one-slide (four steps)", "3-Master_Plan/re_entry/Model.pdf", "Also Model.pptx — Charles updated Sep 2026."),
    ("Act IV", "4.2", "λ knockout?", "Pass A generative sim (G2)", "3-Master_Plan/re_entry/HEROs_and_PASSes/pass_b/PASS_B_generative_lambda_knockout_side_by_side.png", "Repo folder pass_b/; doc 04 = Pass A."),
    ("Act IV", "4.2b", "Empirical talent vs LOO pair", "Pass A empirical (MBB)", "3-Master_Plan/re_entry/HEROs_and_PASSes/pass_a/PASS_A_side_by_side_q16_allt_min20_mg10_9_21_last_ps.png", "Inventory / appendix — same read as 2.2+2.3 on one slide."),
    ("Act IV", "4.3", "Environment ≠ advancement?", "Binding table", "3-Master_Plan/BINDING_Selection_is_its_own_step.md", "Same as slide 0.2."),
    # Act V (optional)
    ("Act V", "5.1", "Why LOO not team mean?", "Usage T̂_j twin + usage LOO", "3-Master_Plan/re_entry/HEROs_and_PASSes/football_sandbox/perf_story/FOOTBALL_usage_Tj_q16_ew16.png", "Pair with usage LOO panel below."),
    ("Act V", "5.1b", "Usage LOO panel", "Perf story row 5", "3-Master_Plan/re_entry/HEROs_and_PASSes/football_sandbox/perf_story/panels/HERO_football_perf_q16_loo_z_diy_usage_overall_slide.png", ""),
    ("Act V", "5.2", "Wrong metric, wrong story?", "PPA LOO vs usage LOO", "3-Master_Plan/re_entry/HEROs_and_PASSes/football_sandbox/data_story/FOOTBALL_PERF_METRIC_STORY.png", "Full six-metric deck; row 4 vs 5 contrast."),
    ("Act V", "5.2b", "PPA total LOO (row 4)", "Talent-peer axis", "3-Master_Plan/re_entry/HEROs_and_PASSes/football_sandbox/perf_story/panels/HERO_football_perf_q16_loo_z_diy_ppa_total_all_slide.png", ""),
    # Act VI
    ("Act VI", "6.1–6.3", "Limits & close", "Text / gap slides", None, "No figures — see flipbook Act VI claims."),
]

# Inventory-only extras (MBB 3×3 building blocks — not flipbook slides; live in inventory deck)
INVENTORY_EXTRA_PANELS: list[tuple[str, str, str]] = [
    (
        "MBB panel 4 · draft-mass ECDF vs Â",
        "3-Master_Plan/re_entry/HEROs_and_PASSes/sports_sandbox/reigning_hero/fhero/BDP_Ai_draft_mass_ecdf_mg10_min20_09_21_allt_ppm_last_ps.png",
        "Band-picking for panels 7–8 (CCT/elite pond) — not naïve-vs-hero pair",
    ),
    ("MBB panel 2 · Â vs T̂_j", "3-Master_Plan/re_entry/HEROs_and_PASSes/sports_sandbox/reigning_hero/basic_data_plots/REIGNING_BDP_Ai_Tj_mg10_min20_09_21_ppm_lastps.png", "Act II context"),
    ("MBB panel 3 · poolq LOO dist", "3-Master_Plan/re_entry/HEROs_and_PASSes/sports_sandbox/reigning_hero/basic_data_plots/REIGNING_BDP_poolq_loo_dist_mg10_min20_09_21_ppm_lastps_nowinsor.png", "Act II context"),
    ("MBB panel 6 · roster size", "3-Master_Plan/re_entry/HEROs_and_PASSes/sports_sandbox/reigning_hero/basic_data_plots/REIGNING_BDP_team_size_mg10_min20_09_21.png", "Act II context"),
    ("MBB panel 7 · CCT", "3-Master_Plan/re_entry/HEROs_and_PASSes/basic_data_plots/CCT_draft_rate_ai_band_poolq_loo_min20_ppm_z2_3_dft.png", "Exploratory · +DFT 11–21"),
    ("MBB panel 8 · elite pond", "3-Master_Plan/re_entry/HEROs_and_PASSes/sports_sandbox/_DISPOSABLE_elite_pond_loo_twin/ELITE_pond_loo_pw4p7_dft_min20_mg10_top7_ppm_11_21.png", "Exploratory · top 7% Â"),
]


def _repo_rel(path: Path) -> str:
    return path.relative_to(REPO).as_posix()


def _iter_army_png_candidates() -> list[Path]:
    """Collect Army figure PNGs from sandbox + talent cox export dirs."""
    seen: set[Path] = set()
    out: list[Path] = []

    def add(path: Path) -> None:
        resolved = path.resolve()
        if resolved in seen or not path.is_file():
            return
        seen.add(resolved)
        out.append(path)

    if ARMY_SANDBOX.exists():
        for ext in ("*.png", "*.jpg", "*.jpeg"):
            for path in ARMY_SANDBOX.rglob(ext):
                add(path)

    talent = REPO / "talent"
    if talent.exists():
        for path in talent.rglob("*.png"):
            parts = {part.lower() for part in path.parts}
            if parts & {"cox_plots", "cox_results", "exports"}:
                add(path)

    return out


def _score_army_candidate(slot: str, path: Path) -> int | None:
    name = path.name.lower()
    path_s = str(path).lower()

    if slot == "hero_cif":
        if "pool_minus_mean" not in name and "poolminusmean" not in name.replace("_", ""):
            return None
        if "bwd" in name:
            return None
        if "partial_effects" in name or "combined_effect" in name:
            return None
        score = 10
        if "fwd" in name:
            score += 5
        if "cif_bars" in name:
            score += 12
        elif "cif" in name:
            score += 6
        if "b8" in name:
            score += 4
        if "_q_" in name or name.endswith("_q_noOER.png"):
            score += 2
    elif slot == "own_tb_cif":
        if "pool_minus_mean" in name:
            return None
        if "bwd" in name:
            return None
        if not any(tok in name for tok in ("tb_r_fwd", "tb_ratio", "z_tb")):
            return None
        score = 10
        if "cif_bars" in name:
            score += 12
        elif "cif" in name:
            score += 6
        if "b8" in name:
            score += 4
    elif slot == "partial_effects":
        if "partial_effects" in name:
            score = 14
            if "pool_minus_mean" in name or "pool" in name:
                score += 4
        elif "combined_effect" in name and "pool_minus_mean" in name:
            score = 12
        elif name.startswith("combined_effect"):
            score = 8
        else:
            return None
    else:
        return None

    if "army_sandbox" in path_s:
        score += 25
    elif "cox_plots" in path_s:
        score += 8
    elif "cox_results" in path_s:
        score += 6
    return score


def discover_army_pngs() -> dict[str, str | None]:
    """Return repo-relative paths for G1 army slots (None if not found)."""
    candidates = _iter_army_png_candidates()
    found: dict[str, str | None] = {k: None for k in ARMY_CANONICAL_FILENAMES}

    # 1) Exact canonical filenames in army_sandbox (any subfolder)
    if ARMY_SANDBOX.exists():
        sandbox_root = ARMY_SANDBOX.resolve()
        by_name: dict[str, Path] = {}
        for path in candidates:
            try:
                path.resolve().relative_to(sandbox_root)
            except ValueError:
                continue
            by_name[path.name] = path
        for slot, names in ARMY_CANONICAL_FILENAMES.items():
            for fname in names:
                if fname in by_name:
                    found[slot] = _repo_rel(by_name[fname])
                    break

    # 2) Pattern match best candidate per slot
    for slot in found:
        if found[slot] is not None:
            continue
        best: tuple[int, float, Path] | None = None
        for path in candidates:
            score = _score_army_candidate(slot, path)
            if score is None:
                continue
            rank = (score, path.stat().st_mtime, path)
            if best is None or rank[:2] > best[:2]:
                best = rank
        if best is not None:
            found[slot] = _repo_rel(best[2])

    return found


def resolve_figures(army: dict[str, str | None]) -> list[tuple[str, str, str, str, str | None, str]]:
    resolved: list[tuple[str, str, str, str, str | None, str]] = []
    for act, slide_id, title, use, repo_path, note in FIGURES:
        slot = FIGURE_ARMY_SLOTS.get((act, slide_id))
        discovered = army.get(slot) if slot else None
        if repo_path is None and discovered:
            repo_path = discovered
            note = f"{note} · auto-discovered".strip(" ·")
        resolved.append((act, slide_id, title, use, repo_path, note))
    return resolved


def _html_img_data_uri(path: Path) -> str:
    b64 = base64.standard_b64encode(path.read_bytes()).decode("ascii")
    mime = "image/png" if path.suffix.lower() == ".png" else "image/jpeg"
    return f"data:{mime};base64,{b64}"


def render_card(act: str, slide: str, title: str, use: str, repo_path: str | None, note: str) -> str:
    path = REPO / repo_path if repo_path else None
    st, st_label = status(path)
    sid = f"{act.replace(' ', '-')}-{slide}".replace(".", "-")
    parts = [
        f'<article class="card {st}" id="{html.escape(sid)}">',
        f'<div class="meta"><span class="act">{html.escape(act)}</span> '
        f'<span class="slide">{html.escape(slide)}</span> '
        f'<span class="badge {st}">{html.escape(st.upper())}</span></div>',
        f"<h2>{html.escape(title)}</h2>",
        f'<p class="use"><strong>Use:</strong> {html.escape(use)}</p>',
    ]
    if repo_path:
        parts.append(f'<p class="path"><code>{html.escape(repo_path)}</code></p>')
    if note:
        parts.append(f'<p class="note">{html.escape(note)}</p>')
    if st == "ok" and path is not None:
        img_rel = rel(path)
        parts.append(f'<a href="{html.escape(img_rel)}" target="_blank" rel="noopener">')
        parts.append(
            f'<img src="{_html_img_data_uri(path)}" alt="{html.escape(title)}" loading="lazy">'
        )
        parts.append("</a>")
        parts.append(f'<p class="path"><code>{html.escape(img_rel)}</code></p>')
    elif st == "pdf" and path is not None:
        parts.append(
            f'<p class="pdf-link"><a href="{html.escape(rel(path))}">Open PDF</a> '
            f"(preview not embedded)</p>"
        )
    elif st == "text" and path is not None and path.exists():
        parts.append(
            f'<p class="pdf-link"><a href="{html.escape(rel(path))}">Open document</a></p>'
        )
    elif st == "missing":
        parts.append('<div class="placeholder">Figure not in repo — see gap note above.</div>')
    elif st == "text":
        parts.append('<div class="placeholder">Text-only slide — no figure file.</div>')
    parts.append("</article>")
    return "\n".join(parts)


def _picture_size_inches(iw: int, ih: int, max_w_in: float, max_h_in: float) -> tuple[float, float]:
    """Fit pixel bitmap into an inch box (display size on slide, not DPI literal)."""
    if iw <= 0 or ih <= 0:
        return 0.0, 0.0
    aspect = iw / ih
    box_aspect = max_w_in / max_h_in
    if aspect >= box_aspect:
        w_in = max_w_in
        h_in = max_w_in / aspect
    else:
        h_in = max_h_in
        w_in = max_h_in * aspect
    return w_in, h_in


def _embed_picture(slide, path: Path, top_in: float, max_w_in: float, max_h_in: float) -> None:
    with Image.open(path) as im:
        iw, ih = im.size
    w_in, h_in = _picture_size_inches(iw, ih, max_w_in, max_h_in)
    if w_in <= 0 or h_in <= 0:
        return
    left_in = MARGIN_IN + (max_w_in - w_in) / 2
    slide.shapes.add_picture(
        str(path),
        Inches(left_in),
        Inches(top_in),
        width=Inches(w_in),
        height=Inches(h_in),
    )


def _add_text_box(slide, left, top, width, height, lines: list[tuple[str, int, bool, RGBColor | None]]):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, (text, size, bold, color) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = text
        p.font.size = Pt(size)
        p.font.bold = bold
        if color is not None:
            p.font.color.rgb = color
        p.space_after = Pt(2)
    return box


def _add_figure_slide(
    prs: Presentation,
    act: str,
    slide_id: str,
    title: str,
    use: str,
    repo_path: str | None,
    note: str,
    *,
    section: str | None = None,
    path_display: str | None = None,
) -> None:
    path = REPO / repo_path if repo_path else None
    st, _ = status(path)
    shown_path = path_display or repo_path or ""

    layout = prs.slide_layouts[6]  # blank
    slide = prs.slides.add_slide(layout)

    if section:
        _add_text_box(
            slide,
            MARGIN,
            Inches(3.0),
            SLIDE_W - 2 * MARGIN,
            Inches(1.5),
            [(section, 32, True, RGBColor(0x33, 0x33, 0x33))],
        )
        return

    header_lines: list[tuple[str, int, bool, RGBColor | None]] = [
        (f"{slide_id} · {title.replace('*', '')}", 20, True, None),
        (f"{act}  ·  {use}", 11, False, RGBColor(0x44, 0x44, 0x44)),
    ]
    _add_text_box(
        slide,
        MARGIN,
        MARGIN,
        SLIDE_W - 2 * MARGIN,
        Inches(HEADER_IN),
        header_lines,
    )

    img_top_in = MARGIN_IN + HEADER_IN + 0.08
    img_max_w_in = SLIDE_W_IN - 2 * MARGIN_IN
    img_max_h_in = SLIDE_H_IN - img_top_in - MARGIN_IN - FOOTER_IN

    if st == "ok" and path is not None:
        _embed_picture(slide, path, img_top_in, img_max_w_in, img_max_h_in)
    elif st == "pdf" and path is not None:
        alt = path.with_suffix(".pptx")
        msg = f"PDF not embedded — open: {repo_path}"
        if alt.exists():
            msg += f"  |  Or: {rel(alt)}"
        _add_text_box(
            slide,
            MARGIN,
            Inches(img_top_in + 1.0),
            SLIDE_W - 2 * MARGIN,
            Inches(1.2),
            [(msg, 14, False, RGBColor(0x15, 0x65, 0xC0))],
        )
    elif st == "text" and path is not None and path.exists():
        _add_text_box(
            slide,
            MARGIN,
            Inches(img_top_in + 1.0),
            SLIDE_W - 2 * MARGIN,
            Inches(1.0),
            [(f"Open document: {repo_path}", 14, False, RGBColor(0x75, 0x75, 0x75))],
        )
    elif st == "missing":
        _add_text_box(
            slide,
            MARGIN,
            Inches(img_top_in + 1.5),
            SLIDE_W - 2 * MARGIN,
            Inches(1.0),
            [("Figure not in repo — see gap note (G1/G2/etc.)", 18, True, RGBColor(0xC6, 0x28, 0x28))],
        )
    else:
        _add_text_box(
            slide,
            MARGIN,
            Inches(img_top_in + 1.5),
            SLIDE_W - 2 * MARGIN,
            Inches(0.8),
            [("Text-only slide — no figure file", 14, False, RGBColor(0x75, 0x75, 0x75))],
        )

    footer_top = Inches(SLIDE_H_IN - MARGIN_IN - FOOTER_IN)
    footer_lines: list[tuple[str, int, bool, RGBColor | None]] = []
    if shown_path:
        footer_lines.append((shown_path, 8, False, RGBColor(0x66, 0x66, 0x66)))
    if note:
        footer_lines.append((note, 8, False, RGBColor(0x66, 0x66, 0x66)))
    status_color = {
        "ok": RGBColor(0x2E, 0x7D, 0x32),
        "missing": RGBColor(0xC6, 0x28, 0x28),
        "pdf": RGBColor(0x15, 0x65, 0xC0),
        "text": RGBColor(0x75, 0x75, 0x75),
    }[st]
    footer_lines.append((st.upper(), 9, True, status_color))
    if footer_lines:
        _add_text_box(slide, MARGIN, footer_top, SLIDE_W - 2 * MARGIN, Inches(FOOTER_IN), footer_lines)


def build_pptx(figures: list[tuple[str, str, str, str, str | None, str]]) -> int:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    title_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_text_box(
        title_slide,
        MARGIN,
        Inches(2.0),
        SLIDE_W - 2 * MARGIN,
        Inches(3.5),
        [
            ("PD30 flipbook — talk skeleton (Acts 0–VI)", 28, True, None),
            ("One slide per flipbook slide ID in _DISPOSABLE_paper_flipbook_PD30.md", 16, False, RGBColor(0x55, 0x55, 0x55)),
            ("Subset of figure inventory — not the full catalog", 14, False, None),
            ("Regenerate: python scripts/build_flipbook_figure_deck.py", 12, False, RGBColor(0x66, 0x66, 0x66)),
        ],
    )

    last_act: str | None = None
    n_images = 0
    for row in figures:
        act, slide_id, title, use, repo_path, note = row
        if act != last_act:
            _add_figure_slide(prs, act, "", "", "", None, "", section=f"— {act} —")
            last_act = act
        _add_figure_slide(prs, act, slide_id, title, use, repo_path, note)
        path = REPO / repo_path if repo_path else None
        if path and status(path)[0] == "ok":
            n_images += 1

    prs.save(str(OUT_PPTX))
    return n_images


def _manifest_image_rows(manifest_path: Path, prefix: str) -> list[tuple]:
    import json

    if not manifest_path.exists():
        return []
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    rows: list[tuple] = []
    for cell in manifest.get("grid") or []:
        if cell.get("type") != "image":
            continue
        p = cell.get("path") or ""
        note = cell.get("note") or ""
        rows.append(
            (
                f"{prefix} · {cell.get('title', 'panel')}",
                p,
                p,
                note or "From reigning 3×3 manifest",
                "",
            )
        )
    return rows


InventoryRow = tuple[str, str | None, str, str, str]


def _inventory_paths(rows: list[InventoryRow]) -> set[str]:
    return {r[1] for r in rows if r[1]}


def _resolve_flipbook_figure_paths(
    figures: list[tuple[str, str, str, str, str | None, str]],
    army: dict[str, str | None],
) -> dict[str, list[str]]:
    """Map repo-relative path → flipbook slide IDs (Act slide_id)."""
    path_to_slides: dict[str, list[str]] = {}
    for act, slide_id, _title, _use, repo_path, _note in figures:
        if repo_path is None:
            slot = FIGURE_ARMY_SLOTS.get((act, slide_id))
            repo_path = army.get(slot) if slot else None
        if not repo_path:
            continue
        ref = f"{act} {slide_id}"
        path_to_slides.setdefault(repo_path, [])
        if ref not in path_to_slides[repo_path]:
            path_to_slides[repo_path].append(ref)
    return path_to_slides


def _tag_flipbook_refs(rows: list[InventoryRow], path_to_slides: dict[str, list[str]]) -> list[InventoryRow]:
    out: list[InventoryRow] = []
    for label, repo_path, path_display, use, note in rows:
        if repo_path and repo_path in path_to_slides:
            fb = ", ".join(path_to_slides[repo_path])
            tag = f"Flipbook: {fb}"
            note = f"{note} · {tag}".strip(" ·") if note else tag
        out.append((label, repo_path, path_display, use, note))
    return out


def _ensure_flipbook_subset(
    primary: list[InventoryRow],
    extended: list[InventoryRow],
    backup: list[InventoryRow],
    figures: list[tuple[str, str, str, str, str | None, str]],
    army: dict[str, str | None],
) -> list[InventoryRow]:
    """Every flipbook figure path must appear in inventory (primary ∪ extended ∪ backup)."""
    known = _inventory_paths(primary) | _inventory_paths(extended) | _inventory_paths(backup)
    path_to_slides = _resolve_flipbook_figure_paths(figures, army)
    added: list[InventoryRow] = []
    for path, slides in sorted(path_to_slides.items()):
        if path in known:
            continue
        slide_ref = slides[0]
        title = next((r[2] for r in figures if _flipbook_row_path(r, army) == path), "Flipbook figure")
        use = next((r[3] for r in figures if _flipbook_row_path(r, army) == path), "In flipbook talk skeleton")
        added.append(
            (
                f"Flipbook {slide_ref} · {title[:50]}",
                path,
                path,
                use,
                f"Auto-added — flipbook slide(s): {', '.join(slides)}",
            )
        )
    return primary + added


def _flipbook_row_path(row: tuple, army: dict[str, str | None]) -> str | None:
    act, slide_id, _t, _u, repo_path, _n = row
    if repo_path:
        return repo_path
    slot = FIGURE_ARMY_SLOTS.get((act, slide_id))
    return army.get(slot) if slot else None


def _build_inventory_extended(existing_paths: set[str]) -> list[InventoryRow]:
    """Germane figures not in primary table — MBB extras, army_sandbox optional exports."""
    rows: list[InventoryRow] = []
    for label, p, tag in INVENTORY_EXTRA_PANELS:
        if p not in existing_paths:
            rows.append((label, p, p, tag, "Inventory extended — not a flipbook slide"))
            existing_paths.add(p)

    if ARMY_SANDBOX.exists():
        for path in sorted(ARMY_SANDBOX.rglob("*.png")):
            if "aws_originals" in path.parts:
                continue
            rel_path = _repo_rel(path)
            if rel_path in existing_paths:
                continue
            rows.append(
                (
                    f"Army sandbox · {path.name}",
                    rel_path,
                    rel_path,
                    "Optional Army export — extended inventory",
                    "Run extract_army_figures_from_520_notebook.py --extras",
                )
            )
            existing_paths.add(rel_path)
    return rows


def _patch_army_inventory_row(
    label: str,
    repo_path: str | None,
    path_display: str,
    use: str,
    note: str,
    army: dict[str, str | None],
) -> tuple:
    slot = INVENTORY_ARMY_SLOTS.get(label)
    discovered = army.get(slot) if slot else None
    if discovered:
        fallback = path_display
        return (
            label,
            discovered,
            discovered,
            use.replace("PNG not in git", "auto-discovered in repo"),
            note or f"Fallback descriptor: {fallback}",
        )
    return label, repo_path, path_display, use, note


def build_inventory_rows(
    army: dict[str, str | None] | None = None,
    figures: list[tuple[str, str, str, str, str | None, str]] | None = None,
) -> tuple[list[InventoryRow], list[InventoryRow], list[InventoryRow]]:
    """Build full figure catalog: primary + extended + backup.

    Policy: every flipbook talk figure ⊆ inventory; inventory may contain more.
    """
    figures = figures or FIGURES
    army = army or discover_army_pngs()
    primary: list[tuple] = [
        _patch_army_inventory_row(
            "Army hero porch (G1)",
            None,
            "520_pipeline_cox_working.ipynb Cell 11 — z_pool_minus_mean_snr_fwd (8-bin promotion CIF)",
            "Lead phenomenon — inverted-U vs peer pool (LOO-style); PNG not in git",
            "G1 — drop in army_sandbox/ or talent/**/cox/cox_plots/",
            army,
        ),
        _patch_army_inventory_row(
            "Army own-TB baseline",
            None,
            "520_pipeline_cox_working.ipynb Cell 11 — z_tb_ratio_fwd_snr (8-bin promotion CIF)",
            "Own top-block share — monotone (pairs with hero)",
            "Cell 11 export — auto-scan cox_plots / army_sandbox",
            army,
        ),
        _patch_army_inventory_row(
            "Army Cox / partial effects",
            None,
            "520 Cell 12 + 12.6A — quadratic z_pool_minus_mean_snr_fwd_sq",
            "Formal curvature check vs CIF bars",
            "Cell 12 / 12.6A export — auto-scan cox_results / army_sandbox",
            army,
        ),
        (
            "Army pipeline / talk track",
            "talent/documents/520_PIPELINE_COX_OVERVIEW.md",
            "talent/documents/520_PIPELINE_COX_OVERVIEW.md · Presentation_Interpretation_Run1_Slides_5-20.md",
            "502→512→520; Run 1 CS+CSS YG ~2002–2013",
            "Notebook: talent/talent_pipeline/520_pipeline_cox_working.ipynb",
        ),
        (
            "MBB 3×3",
            "3-Master_Plan/re_entry/HEROs_and_PASSes/sports_sandbox/data_story/MBB_DATA_STORY_reigning_3x3.png",
            "3-Master_Plan/re_entry/HEROs_and_PASSes/sports_sandbox/data_story/MBB_DATA_STORY_reigning_3x3.png",
            "Replication porch while CAC out; panel 9 = reigning hero",
            "",
        ),
        (
            "MBB highlights (talk track doc)",
            "3-Master_Plan/re_entry/HEROs_and_PASSes/sports_sandbox/data_story/MBB_DATA_STORY_plot_highlights.md",
            "…/sports_sandbox/data_story/MBB_DATA_STORY_plot_highlights.md",
            "Talk track for each panel — see following manifest panels",
            "",
        ),
        (
            "Tenure 3×3",
            "3-Master_Plan/re_entry/HEROs_and_PASSes/tenure_sandbox/data_story/TENURE_DATA_STORY_pd29_3x3.png",
            "3-Master_Plan/re_entry/HEROs_and_PASSes/tenure_sandbox/data_story/TENURE_DATA_STORY_pd29_3x3.png",
            "PD29 decision cohort porch",
            "",
        ),
        (
            "Tenure highlights (talk track doc)",
            "3-Master_Plan/re_entry/HEROs_and_PASSes/tenure_sandbox/data_story/TENURE_DATA_STORY_plot_highlights.md",
            "…/tenure_sandbox/data_story/TENURE_DATA_STORY_plot_highlights.md",
            "Panel 9 β₂ ≈ −0.017 concave — see following manifest panels",
            "",
        ),
        (
            "Tenure perf p1",
            "3-Master_Plan/re_entry/HEROs_and_PASSes/tenure_sandbox/data_story/TENURE_PERF_METRIC_STORY.png",
            "…/tenure_sandbox/data_story/TENURE_PERF_METRIC_STORY.png",
            "Career rate · cum pubs · annum pubs",
            "",
        ),
        (
            "Tenure perf p2",
            "3-Master_Plan/re_entry/HEROs_and_PASSes/tenure_sandbox/data_story/TENURE_PERF_METRIC_STORY_p2.png",
            "…/tenure_sandbox/data_story/TENURE_PERF_METRIC_STORY_p2.png",
            "10-bin sensitivity",
            "",
        ),
        (
            "Model one-slide",
            "3-Master_Plan/re_entry/Model.pdf",
            "3-Master_Plan/re_entry/Model.pptx · Model.pdf",
            "Assign → score → select; unified S_i",
            "Open Model.pptx for the actual slide",
        ),
        (
            "Pass A sim (λ knockout) (G2)",
            "3-Master_Plan/re_entry/HEROs_and_PASSes/pass_b/PASS_B_generative_lambda_knockout_side_by_side.png",
            "…/pass_b/PASS_B_generative_lambda_knockout_side_by_side.png",
            "Talent-only vs congestion-in-score; doc 04 = Pass A",
            "",
        ),
        (
            "Pass A empirical pair",
            "3-Master_Plan/re_entry/HEROs_and_PASSes/pass_a/PASS_A_empirical_talent_vs_roster_side_by_side_16quantile_winsor0199_min20_q16.png",
            "…/pass_a/PASS_A_empirical_talent_vs_roster_side_by_side_16quantile_winsor0199_min20_q16.png",
            "Â monotone vs poolq LOO hero — not the sim slide",
            "",
        ),
        (
            "Football usage twin",
            "3-Master_Plan/re_entry/HEROs_and_PASSes/football_sandbox/perf_story/FOOTBALL_usage_Tj_q16_ew16.png",
            "…/football_sandbox/perf_story/FOOTBALL_usage_Tj_q16_ew16.png",
            "T̂_j flat vs LOO steep — why LOO",
            "",
        ),
        (
            "Football usage LOO",
            "3-Master_Plan/re_entry/HEROs_and_PASSes/football_sandbox/perf_story/panels/HERO_football_perf_q16_loo_z_diy_usage_overall_slide.png",
            "…/football_sandbox/perf_story/panels/HERO_football_perf_q16_loo_z_diy_usage_overall_slide.png",
            "Perf story row 5",
            "",
        ),
        (
            "Football PPA LOO",
            "3-Master_Plan/re_entry/HEROs_and_PASSes/football_sandbox/perf_story/panels/HERO_football_perf_q16_loo_z_diy_ppa_total_all_slide.png",
            "…/football_sandbox/perf_story/panels/HERO_football_perf_q16_loo_z_diy_ppa_total_all_slide.png",
            "Perf story row 4 — talent-peer contrast",
            "",
        ),
        (
            "Football perf deck",
            "3-Master_Plan/re_entry/HEROs_and_PASSes/football_sandbox/data_story/FOOTBALL_PERF_METRIC_STORY.png",
            "…/football_sandbox/data_story/FOOTBALL_PERF_METRIC_STORY.png",
            "Six-metric screen (Act V optional)",
            "",
        ),
        (
            "Binding doc",
            "3-Master_Plan/BINDING_Selection_is_its_own_step.md",
            "3-Master_Plan/BINDING_Selection_is_its_own_step.md",
            "Score ≠ select slide",
            "",
        ),
    ]

    # Insert MBB manifest panels after MBB highlights row
    mbb_manifest = REPO / "3-Master_Plan/re_entry/HEROs_and_PASSes/sports_sandbox/data_story/mbb_reigning_3x3_manifest.json"
    mbb_panels = _manifest_image_rows(mbb_manifest, "MBB manifest")
    insert_at = next(i for i, r in enumerate(primary) if r[0].startswith("MBB highlights"))
    primary[insert_at + 1 : insert_at + 1] = mbb_panels

    tenure_manifest = REPO / "3-Master_Plan/re_entry/HEROs_and_PASSes/tenure_sandbox/data_story/tenure_pd29_3x3_manifest.json"
    tenure_panels = _manifest_image_rows(tenure_manifest, "Tenure manifest")
    insert_at = next(i for i, r in enumerate(primary) if r[0].startswith("Tenure highlights"))
    primary[insert_at + 1 : insert_at + 1] = tenure_panels

    backup: list[InventoryRow] = [
        (
            "Big Fish · Football 3×3 (backup)",
            "3-Master_Plan/re_entry/HEROs_and_PASSes/football_sandbox/data_story/FOOTBALL_DATA_STORY_3x3.png",
            "…/football_sandbox/data_story/FOOTBALL_DATA_STORY_3x3.png",
            "PD30 back burner — not a fourth anchor leg",
            "",
        ),
        (
            "Big Fish · Legends 3×3 (backup)",
            "3-Master_Plan/re_entry/HEROs_and_PASSes/legends_sandbox/data_story/LEGENDS_DATA_STORY_3x3.png",
            "…/legends_sandbox/data_story/LEGENDS_DATA_STORY_3x3.png",
            "PD30 back burner — weak porch",
            "",
        ),
    ]

    path_to_slides = _resolve_flipbook_figure_paths(figures, army)
    primary = _tag_flipbook_refs(primary, path_to_slides)
    extended = _build_inventory_extended(_inventory_paths(primary))
    extended = _tag_flipbook_refs(extended, path_to_slides)
    primary = _ensure_flipbook_subset(primary, extended, backup, figures, army)
    return primary, extended, backup


def _inventory_card(label: str, repo_path: str | None, path_display: str, use: str, note: str) -> str:
    path = REPO / repo_path if repo_path else None
    st, _ = status(path)
    sid = "inv-" + label.lower().replace(" ", "-").replace("(", "").replace(")", "")[:40]
    parts = [
        f'<article class="card {st}" id="{html.escape(sid)}">',
        f'<div class="meta"><span class="act">Figure inventory</span> '
        f'<span class="badge {st}">{html.escape(st.upper())}</span></div>',
        f"<h2>{html.escape(label)}</h2>",
        f'<p class="use"><strong>Use in talk:</strong> {html.escape(use)}</p>',
        f'<p class="path"><code>{html.escape(path_display or "")}</code></p>',
    ]
    if note:
        parts.append(f'<p class="note">{html.escape(note)}</p>')
    if st == "ok" and path is not None and path.suffix.lower() in {".png", ".jpg", ".jpeg"}:
        parts.append(f'<img src="{_html_img_data_uri(path)}" alt="{html.escape(label)}" loading="lazy">')
    elif st == "missing":
        parts.append('<div class="placeholder">No PNG in repo — path is descriptive / notebook export pending.</div>')
    elif st == "text":
        parts.append('<div class="placeholder">Document-only inventory row — open path in editor.</div>')
    elif st == "pdf":
        parts.append('<div class="placeholder">PDF — open Model.pptx / Model.pdf locally.</div>')
    parts.append("</article>")
    return "\n".join(parts)


def build_inventory_pptx(
    primary: list[InventoryRow],
    extended: list[InventoryRow],
    backup: list[InventoryRow],
) -> int:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    title = prs.slides.add_slide(prs.slide_layouts[6])
    _add_text_box(
        title,
        MARGIN,
        Inches(2.0),
        SLIDE_W - 2 * MARGIN,
        Inches(3.5),
        [
            ("Figure inventory — full catalog (PD30)", 28, True, None),
            ("All germane figures · flipbook talk figures are a subset", 16, False, RGBColor(0x55, 0x55, 0x55)),
            ("Companion to _DISPOSABLE_paper_flipbook_PD30.md § Figure inventory", 14, False, None),
        ],
    )

    n_images = 0

    def _add_rows(rows: list[InventoryRow], section_label: str, id_prefix: str) -> None:
        nonlocal n_images
        if not rows:
            return
        _add_figure_slide(prs, "", "", "", "", None, "", section=section_label)
        for idx, (label, repo_path, path_display, use, note) in enumerate(rows, start=1):
            sid = f"{id_prefix}{idx:02d}" if id_prefix else f"{idx:02d}"
            _add_figure_slide(
                prs,
                "Figure inventory",
                sid,
                label,
                use,
                repo_path,
                note,
                path_display=path_display,
            )
            path = REPO / repo_path if repo_path else None
            if path and status(path)[0] == "ok" and path.suffix.lower() in {".png", ".jpg", ".jpeg"}:
                n_images += 1

    _add_rows(primary, "— Primary + manifest panels —", "")
    _add_rows(extended, "— Extended (germane, not in flipbook talk) —", "X")
    _add_figure_slide(prs, "", "", "", "", None, "", section="— Big Fish backup (PD30) —")
    for label, repo_path, path_display, use, note in backup:
        _add_figure_slide(
            prs,
            "Backup",
            "—",
            label,
            use,
            repo_path,
            note,
            path_display=path_display,
        )
        path = REPO / repo_path if repo_path else None
        if path and status(path)[0] == "ok":
            n_images += 1

    prs.save(str(OUT_INV_PPTX))
    return n_images


def build_inventory_html(
    primary: list[InventoryRow],
    extended: list[InventoryRow],
    backup: list[InventoryRow],
) -> None:
    cards = [_inventory_card(*row) for row in primary]
    if extended:
        cards.append('<section class="extra-block"><h2>Extended inventory (germane, not flipbook talk)</h2></section>')
        cards.extend(_inventory_card(*row) for row in extended)
    cards.append('<section class="extra-block"><h2>Big Fish backup (PD30)</h2></section>')
    cards.extend(_inventory_card(*row) for row in backup)
    toc = "".join(
        f'<li><a href="#inv-{html.escape(row[0].lower().replace(" ", "-")[:40])}">{html.escape(row[0])}</a></li>'
        for row in primary + extended + backup
    )
    doc = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<title>PD30 Figure Inventory — visual catalog</title>
<style>
  body {{ font-family: system-ui, sans-serif; max-width: 1100px; margin: 0 auto; padding: 1.5rem; background: #f6f6f8; }}
  header {{ background: #fff; padding: 1.25rem; border-radius: 8px; margin-bottom: 1rem; border: 1px solid #ddd; }}
  .card {{ background: #fff; border-radius: 8px; padding: 1rem; margin-bottom: 1rem; border: 1px solid #ddd; }}
  .card.ok {{ border-left: 4px solid #2e7d32; }}
  .card.missing {{ border-left: 4px solid #c62828; }}
  .card.text, .card.pdf {{ border-left: 4px solid #757575; }}
  img {{ max-width: 100%; height: auto; margin-top: 0.75rem; border: 1px solid #ccc; }}
  .placeholder {{ padding: 1.5rem; background: #fafafa; border: 2px dashed #ccc; color: #666; text-align: center; }}
  code {{ font-size: 0.85rem; word-break: break-all; }}
  nav ul {{ columns: 2; font-size: 0.9rem; }}
</style>
</head>
<body>
<header>
  <h1>Figure inventory — visual catalog</h1>
  <p>Full figure catalog for PD30 — <strong>flipbook talk figures ⊆ this inventory</strong>, not vice versa.
     Primary table + manifest panels + extended germane figures + Big Fish backup.</p>
</header>
<nav><h2>Contents</h2><ul>{toc}</ul></nav>
<main>{"".join(cards)}</main>
</body>
</html>"""
    OUT_INV_HTML.write_text(doc, encoding="utf-8")


def main() -> None:
    army = discover_army_pngs()
    figures = resolve_figures(army)

    cards = [render_card(*row) for row in figures]

    toc_html = []
    last_act = None
    for row in figures:
        act, slide, title = row[0], row[1], row[2]
        if act != last_act:
            if last_act is not None:
                toc_html.append("</ul>")
            toc_html.append(f"<h3>{html.escape(act)}</h3><ul>")
            last_act = act
        sid = f"{act.replace(' ', '-')}-{slide}".replace(".", "-")
        toc_html.append(f'<li><a href="#{sid}">{html.escape(slide)} · {html.escape(title)}</a></li>')
    toc_html.append("</ul>")

    doc = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>PD30 Flipbook — Annotated Figure Deck</title>
<style>
  :root {{ font-family: system-ui, -apple-system, sans-serif; line-height: 1.45; color: #1a1a1a; }}
  body {{ max-width: 1100px; margin: 0 auto; padding: 1.5rem; background: #f6f6f8; }}
  header {{ background: #fff; padding: 1.25rem 1.5rem; border-radius: 8px; margin-bottom: 1.5rem; border: 1px solid #ddd; }}
  header h1 {{ margin: 0 0 0.5rem; font-size: 1.35rem; }}
  header p {{ margin: 0.35rem 0; color: #444; font-size: 0.95rem; }}
  nav.toc {{ background: #fff; padding: 1rem 1.25rem; border-radius: 8px; margin-bottom: 1.5rem; border: 1px solid #ddd; }}
  nav.toc h2 {{ margin-top: 0; font-size: 1.1rem; }}
  nav.toc h3 {{ margin: 0.75rem 0 0.25rem; font-size: 0.95rem; color: #555; }}
  nav.toc ul {{ margin: 0 0 0.5rem; padding-left: 1.25rem; }}
  nav.toc a {{ color: #0645ad; text-decoration: none; font-size: 0.9rem; }}
  nav.toc a:hover {{ text-decoration: underline; }}
  .card {{ background: #fff; border-radius: 8px; padding: 1rem 1.25rem 1.25rem; margin-bottom: 1.25rem; border: 1px solid #ddd; }}
  .card.missing {{ border-left: 4px solid #c62828; }}
  .card.ok {{ border-left: 4px solid #2e7d32; }}
  .card.pdf {{ border-left: 4px solid #1565c0; }}
  .card.text {{ border-left: 4px solid #757575; }}
  .meta {{ font-size: 0.85rem; margin-bottom: 0.5rem; }}
  .act {{ font-weight: 600; color: #333; }}
  .slide {{ background: #eee; padding: 0.1rem 0.45rem; border-radius: 4px; margin-left: 0.35rem; }}
  .badge {{ float: right; font-size: 0.7rem; font-weight: 700; padding: 0.15rem 0.5rem; border-radius: 4px; }}
  .badge.ok {{ background: #e8f5e9; color: #2e7d32; }}
  .badge.missing {{ background: #ffebee; color: #c62828; }}
  .badge.pdf {{ background: #e3f2fd; color: #1565c0; }}
  .badge.text {{ background: #f5f5f5; color: #616161; }}
  .card h2 {{ margin: 0.25rem 0 0.5rem; font-size: 1.15rem; }}
  .use {{ margin: 0.4rem 0; }}
  .path code, .note {{ font-size: 0.82rem; color: #555; word-break: break-all; }}
  .note {{ margin-top: 0.35rem; }}
  img {{ max-width: 100%; height: auto; margin-top: 0.75rem; border: 1px solid #ccc; border-radius: 4px; display: block; }}
  .placeholder {{ margin-top: 0.75rem; padding: 2rem; background: #fafafa; border: 2px dashed #ccc; text-align: center; color: #888; border-radius: 4px; }}
  section.extra-block h2 {{ font-size: 1.2rem; margin: 2rem 0 1rem; }}
  @media print {{
    body {{ background: #fff; }}
    .card {{ break-inside: avoid; page-break-inside: avoid; }}
    nav.toc {{ break-after: page; }}
  }}
</style>
</head>
<body>
<header>
  <h1>PD30 flipbook — talk skeleton (Acts 0–VI)</h1>
  <p>Companion to <code>_DISPOSABLE_paper_flipbook_PD30.md</code> § Slide deck — one card per flipbook slide ID.</p>
  <p>Not the full figure catalog — see <code>_DISPOSABLE_paper_flipbook_PD30_FIGURE_INVENTORY.pptx</code> for all germane figures.</p>
  <p><strong>Green</strong> = PNG in repo · <strong>Red</strong> = missing · <strong>Blue</strong> = PDF · <strong>Gray</strong> = text only.</p>
  <p>Generated by <code>scripts/build_flipbook_figure_deck.py</code> — re-run after new PNG exports.</p>
</header>
<nav class="toc">
  <h2>Jump to slide</h2>
  {"".join(toc_html)}
</nav>
<main>
{"".join(cards)}
</main>
</body>
</html>
"""
    OUT_HTML.write_text(doc, encoding="utf-8")
    n_pptx = build_pptx(figures)
    primary_inv, extended_inv, backup_inv = build_inventory_rows(army, figures)
    n_inv = build_inventory_pptx(primary_inv, extended_inv, backup_inv)
    build_inventory_html(primary_inv, extended_inv, backup_inv)
    fb_paths = set(_resolve_flipbook_figure_paths(figures, army))
    inv_paths = _inventory_paths(primary_inv) | _inventory_paths(extended_inv) | _inventory_paths(backup_inv)
    missing_fb = sorted(fb_paths - inv_paths)
    print(f"Wrote {OUT_HTML}")
    print(f"Wrote {OUT_PPTX}")
    print(f"Wrote {OUT_INV_HTML}")
    print(f"Wrote {OUT_INV_PPTX}")
    print(f"Act deck: {len(figures)} cards · {n_pptx} PPTX images")
    print(
        f"Inventory deck: {len(primary_inv)} primary + {len(extended_inv)} extended + "
        f"{len(backup_inv)} backup · {n_inv} PPTX images"
    )
    if missing_fb:
        print("WARNING: flipbook paths missing from inventory:", missing_fb, file=sys.stderr)
    n_scanned = len(_iter_army_png_candidates())
    print(f"Army scan: {n_scanned} PNG candidate(s) under army_sandbox + talent/cox_*")
    for slot, repo_rel in army.items():
        status_label = repo_rel if repo_rel else "(not found — export to army_sandbox/ or cox_plots/)"
        print(f"  {slot}: {status_label}")


if __name__ == "__main__":
    main()
